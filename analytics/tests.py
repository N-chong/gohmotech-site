from datetime import datetime, timedelta, timezone as dt_timezone
from decimal import Decimal

from django.contrib.auth.models import User
from django.test import SimpleTestCase, TestCase
from django.urls import reverse
from django.utils import timezone
from django.http import QueryDict

from analytics import report_data
from analytics.exports import export_csv, export_xlsx
from analytics.models import Report
from iot.models import Device, Goat, SensorReading
from iot.goat_models import GoatWeightMeasurement
from marketplace.auth import BUYER_GROUP_NAME
from marketplace.auth import FARM_OWNER_GROUP_NAME
from marketplace.models import (
    Conversation,
    MarketplaceListing,
    MarketplaceReport,
    Reservation,
    SellerProfile,
)
from security.models import Notification, PersonDetection
from django.contrib.auth.models import Group


class MonthlyAggregationTests(SimpleTestCase):
    def test_monthly_counts_handle_timezone_aware_dates(self):
        now = timezone.make_aware(datetime(2024, 2, 15, 12, 0, 0))
        values = [
            timezone.make_aware(datetime(2024, 1, 5, 10, 0, 0)),
            timezone.make_aware(datetime(2023, 12, 20, 8, 0, 0)),
        ]

        labels, counts = report_data._monthly_counts_for_last_year(values, now=now)

        self.assertEqual(labels[0], "Mar 2023")
        self.assertEqual(labels[-1], "Feb 2024")
        self.assertEqual(counts[10], 1)
        self.assertEqual(counts[11], 0)
        self.assertEqual(counts[9], 1)
        self.assertEqual(len(labels), 12)


class ReportTimezoneTests(TestCase):
    def test_environment_report_converts_stored_utc_to_manila(self):
        device = Device.objects.create(
            device_id="report-time-device",
            name="Report Time Device",
            device_type="environmental_sensor",
        )
        SensorReading.objects.create(
            device=device,
            temperature=30,
            timestamp=datetime(2026, 10, 2, 9, 38, tzinfo=dt_timezone.utc),
        )

        data = report_data.build_environmental(days=3650)

        self.assertEqual(data["rows"][0][0], "2026-10-02 17:38")

    def test_shared_formatter_does_not_add_eight_hours_twice(self):
        manila_value = timezone.make_aware(datetime(2026, 10, 2, 17, 38))
        self.assertEqual(
            report_data._format_local_datetime(manila_value),
            "2026-10-02 17:38",
        )


class ReportingPeriodAndAccuracyTests(TestCase):
    def setUp(self):
        self.device = Device.objects.create(
            device_id="report-environment-device",
            name="Barn Environment",
            device_type="environmental_sensor",
        )

    def test_environment_report_never_exposes_unsupported_soil_moisture(self):
        SensorReading.objects.create(
            device=self.device,
            temperature=28.4,
            humidity=71,
            soil_moisture=42,
        )

        data = report_data.build_environmental(days=1)
        labels = [card["label"] for card in data["summary"]]

        self.assertNotIn("Soil Moisture (%)", data["columns"])
        self.assertFalse(any("soil" in label.lower() for label in labels))
        self.assertEqual(len(data["columns"]), len(data["rows"][0]))
        self.assertNotIn("soil", export_csv(data).content.decode("utf-8").lower())

        from openpyxl import load_workbook
        from io import BytesIO
        workbook = load_workbook(BytesIO(export_xlsx(data).content), read_only=True)
        sheet_text = " ".join(
            str(cell.value or "") for row in workbook.active.iter_rows() for cell in row
        ).lower()
        self.assertNotIn("soil", sheet_text)

    def test_custom_period_is_inclusive_of_the_selected_end_date(self):
        today = timezone.localdate()
        get = QueryDict(mutable=True)
        get.update({
            "period": "custom",
            "start_date": today.isoformat(),
            "end_date": today.isoformat(),
        })

        window = report_data.resolve_period(get)

        self.assertEqual(window["start_date"], today)
        self.assertEqual(window["end_date"], today)
        self.assertEqual(window["end"] - window["start"], timedelta(days=1))

    def test_invalid_custom_period_is_rejected(self):
        get = QueryDict("period=custom&start_date=2026-10-04&end_date=2026-10-01")
        with self.assertRaisesMessage(ValueError, "start date"):
            report_data.resolve_period(get)

    def test_security_report_uses_stored_person_detections(self):
        PersonDetection.objects.create(
            source_label="Barn Camera",
            confidence=0.87,
            detection_type="person",
        )
        Notification.objects.create(
            title="Person detected",
            notification_type="person_detection",
            severity="high",
        )

        data = report_data.build_security(days=1)
        summary = {card["label"]: card["value"] for card in data["summary"]}

        self.assertEqual(summary["Person Detections"], 1)
        self.assertEqual(summary["Security Alerts"], 1)
        self.assertEqual(data["rows"][0][1], "Barn Camera")
        self.assertIn("does not identify people", data["note"])

    def test_weight_report_uses_confirmed_measurement_history(self):
        goat = Goat.objects.create(
            goat_id="WEIGHT-001", breed="native", gender="female"
        )
        GoatWeightMeasurement.objects.create(
            goat=goat, weight_kg=Decimal("34.50"), source="load_cell"
        )

        data = report_data.build_weight(days=1)
        summary = {card["label"]: card["value"] for card in data["summary"]}

        self.assertEqual(summary["Measurements"], 1)
        self.assertEqual(summary["Average Weight"], "34.5 kg")
        self.assertEqual(data["rows"][0][1], "WEIGHT-001")

    def test_print_report_has_professional_title_and_period(self):
        owner = User.objects.create_user(
            "reportowner", password="testpass123", is_staff=True
        )
        self.client.force_login(owner)

        response = self.client.get(
            reverse("analytics:report_environmental_print"), {"period": "today"}
        )

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "GOHMOTECH FARM MONITORING REPORT")
        self.assertContains(response, "Reporting Period:")
        self.assertNotContains(response, "Soil Moisture")

    def test_report_center_and_cross_module_pdf_render_with_empty_categories(self):
        owner = User.objects.create_user(
            "overviewowner", password="testpass123", is_staff=True
        )
        self.client.force_login(owner)

        center = self.client.get(reverse("analytics:dashboard"), {"period": "today"})
        pdf_view = self.client.get(
            reverse("analytics:farm_report_print"), {"period": "today"}
        )

        self.assertEqual(center.status_code, 200)
        self.assertContains(center, "Generate Farm PDF")
        self.assertEqual(pdf_view.status_code, 200)
        self.assertContains(pdf_view, "Livestock Summary")
        self.assertContains(pdf_view, "No security detections were recorded")
        self.assertNotContains(pdf_view, "Soil Moisture")

    def test_all_registered_report_pages_render(self):
        owner = User.objects.create_user(
            "allreportsowner", password="testpass123", is_staff=True
        )
        self.client.force_login(owner)

        for report_type in report_data.REPORTS:
            with self.subTest(report_type=report_type):
                response = self.client.get(
                    reverse(f"analytics:report_{report_type}"), {"period": "today"}
                )
                self.assertEqual(response.status_code, 200)

    def test_saved_report_api_does_not_expose_another_owners_metadata(self):
        owner_group = Group.objects.get_or_create(name=FARM_OWNER_GROUP_NAME)[0]
        owner = User.objects.create_user("apiowner", password="testpass123")
        other = User.objects.create_user("otherowner", password="testpass123")
        owner.groups.add(owner_group)
        other.groups.add(owner_group)
        Report.objects.create(
            title="Mine", report_type="daily", start_date="2026-10-01",
            end_date="2026-10-01", generated_by=owner,
        )
        Report.objects.create(
            title="Other farm", report_type="daily", start_date="2026-10-01",
            end_date="2026-10-01", generated_by=other,
        )
        self.client.force_login(owner)

        response = self.client.get(reverse("analytics:report-list"))

        self.assertEqual(response.status_code, 200)
        titles = [item["title"] for item in response.json()["results"]]
        self.assertEqual(titles, ["Mine"])


class MarketplaceReportTests(TestCase):
    def setUp(self):
        self.staff = User.objects.create_user(
            "reportadmin", password="testpass123", is_staff=True
        )
        self.buyer = User.objects.create_user("reportbuyer", password="testpass123")
        self.buyer.groups.add(Group.objects.get_or_create(name=BUYER_GROUP_NAME)[0])
        SellerProfile.objects.create(
            user=self.buyer,
            farm_name="Report Farm",
            municipality="San Jose",
            province="Nueva Ecija",
            contact_number="09171234567",
            farm_description="Analytics test seller.",
            status=SellerProfile.APPROVED,
        )
        self.goat = Goat.objects.create(
            goat_id="REPORT-GOAT-001",
            breed="native",
            gender="female",
            status="sold",
            is_active=False,
        )
        self.listing = MarketplaceListing.objects.create(
            goat=self.goat,
            seller=self.staff,
            price=Decimal("15000.00"),
            sales_description="Report test listing",
            status=MarketplaceListing.SOLD,
            sold_at=timezone.now(),
        )
        self.conversation = Conversation.objects.create(
            listing=self.listing, buyer=self.buyer
        )
        self.reservation = Reservation.objects.create(
            listing=self.listing,
            buyer=self.buyer,
            conversation=self.conversation,
            agreed_price=Decimal("15000.00"),
            status=Reservation.COMPLETED,
            pickup_status=Reservation.PICKUP_CONFIRMED,
            pickup_datetime=timezone.now(),
            completed_at=timezone.now(),
            completed_by=self.staff,
        )

    def test_marketplace_builder_uses_completed_sale_revenue(self):
        data = report_data.build_marketplace(days=30)
        summary = {card["label"]: card["value"] for card in data["summary"]}
        self.assertEqual(summary["Completed Sales"], 1)
        self.assertEqual(summary["Sales Revenue"], "PHP 15,000.00")
        self.assertEqual(summary["Registered Buyers"], 1)
        self.assertEqual(summary["Approved Sellers"], 1)
        self.assertEqual(len(data["rows"]), 1)
        self.assertEqual(data["rows"][0][1], "REPORT-GOAT-001")
        chart_ids = {chart["id"] for chart in data["charts"]}
        self.assertTrue(
            {"seller_status", "listed_breeds", "listing_locations"}.issubset(chart_ids)
        )

    def test_marketplace_builder_counts_open_reports(self):
        MarketplaceReport.objects.create(
            reporter=self.buyer,
            listing=self.listing,
            reason=MarketplaceReport.INCORRECT,
        )
        data = report_data.build_marketplace(days=30)
        summary = {card["label"]: card["value"] for card in data["summary"]}
        self.assertEqual(summary["Open Reports"], 1)

    def test_marketplace_report_is_in_report_center_and_renders(self):
        self.client.force_login(self.staff)
        center = self.client.get(reverse("analytics:dashboard"))
        self.assertEqual(center.status_code, 200)
        self.assertContains(center, "Marketplace Report")

        detail = self.client.get(reverse("analytics:report_marketplace"))
        self.assertEqual(detail.status_code, 200)
        self.assertContains(detail, "PHP 15,000.00")
        self.assertContains(detail, "REPORT-GOAT-001")

    def test_marketplace_csv_export_uses_same_report_data(self):
        self.client.force_login(self.staff)
        response = self.client.get(
            reverse("analytics:report_marketplace_export"),
            {"format": "csv", "days": "30", "status": "completed"},
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"], "text/csv")
        content = response.content.decode("utf-8")
        self.assertIn("Marketplace Report", content)
        self.assertIn("REPORT-GOAT-001", content)
        self.assertIn("PHP 15,000.00", content)
